"""Build the 10-slide presentation for the Insurance Claims project."""
from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_slide(title, bullets):
    slide = prs.slides.add_slide(BLANK)
    tx = slide.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12.1), Inches(0.9))
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    box = slide.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(12.1), Inches(5.5))
    tf = box.text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "•  " + b
        p.font.size = Pt(22)
        p.space_after = Pt(14)
    return slide


add_slide("Insurance Policy and Claims Management System", [
    "A Spring Boot backend + React frontend for insurance customers, policies, premiums, and claims",
    "Portfolio project by Ramesh Poluri — fictional demo data, no employer claims",
    "Java 17 · Spring Boot 3.2 · Spring Data JPA · Spring Security · MySQL 8 · React + Vite · JUnit 5 + Mockito",
])
add_slide("What the system does", [
    "Customer profiles with unique email addresses",
    "Policy creation, management, and status updates",
    "Premium calculation from policy attributes",
    "Full claim lifecycle: registered → review → approved → settled",
    "Event-driven notifications on every policy and claim change",
    "React web UI for customers, policies, quotes, and claims",
])
add_slide("Architecture", [
    "React frontend (Vite) → Spring Boot controllers → services → repositories → MySQL 8",
    "Centralized exception handling: one clean error body everywhere",
    "EventPublisher interface: local publisher by default, HTTP publisher opt-in",
    "Spring Security: HTTP Basic auth on every API endpoint, CORS enabled for the frontend dev server",
    "Tests run on H2 (in-memory); runtime uses MySQL via docker-compose",
])
add_slide("Domain model", [
    "Customer: name, email, phone, date of birth, address",
    "Policy: policy number, type, coverage, premium, effective and renewal dates, status",
    "Claim: claim number, policy link, amount, incident date, status, settled amount",
    "Policy types: HEALTH, AUTO, HOME, LIFE, TRAVEL",
])
add_slide("Premium calculation", [
    "premium = coverage × base rate × age factor",
    "Base rates: HEALTH 5%, AUTO 4%, HOME 3%, LIFE 2.5%, TRAVEL 6%",
    "Age factor: under 25 → 1.10, 60 and over → 1.20, otherwise 1.00",
    "Example: HEALTH policy, 100,000 coverage, age 36 → 5,000.00",
    "Available as a no-commit quote endpoint before creating a policy",
])
add_slide("Claim lifecycle", [
    "REGISTERED → UNDER_REVIEW → APPROVED → SETTLED",
    "REJECTED is a terminal branch from REGISTERED or UNDER_REVIEW",
    "Illegal transitions are rejected — a claim can never jump states or settle twice",
    "Settlement pays min(claim amount, remaining coverage on the policy)",
])
add_slide("Validation and business rules", [
    "Customers must be 18 or older to hold a policy",
    "Renewal date must be after the effective date",
    "Claims only against ACTIVE policies, inside the policy period",
    "Claim amounts cannot exceed the coverage still unused on the policy",
    "Bean Validation on every request plus service-level rule checks",
])
add_slide("Events and messaging", [
    "Every creation and status change publishes a DomainEvent",
    "Local publisher (default): events logged and kept in memory — no broker needed",
    "HTTP publisher (opt-in): events POSTed as JSON to a configured endpoint",
    "Delivery failures are logged, never fatal to the business operation",
])
add_slide("Security and testing", [
    "HTTP Basic auth with a local demo user (demo / demo123) — demo only, no OAuth provider",
    "Unauthenticated requests get a clean 401",
    "JUnit 5 + Mockito: premium math, service rules, lifecycle transitions",
    "MockMvc integration tests: full customer → policy → claim → settle flow",
])
add_slide("Run it and what's next", [
    "Run: docker compose up --build (MySQL + backend) and cd frontend && npm install && npm run dev",
    "Or backend only: mvn spring-boot:run with a local MySQL — then try the Postman collection",
    "CI: GitHub Actions runs the test suite (tests use H2, no MySQL server needed)",
    "GitHub: github.com/rpoluri9/insurance-policy-claims",
    "Honest scope: local-first demo — real messaging and identity providers are opt-in configuration",
])

prs.save("/home/hatch/workspace/your_files/insurance-claims-presentation.pptx")
print("presentation saved")
